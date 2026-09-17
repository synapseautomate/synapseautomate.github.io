#!/usr/bin/env python3
"""Safe local ingestion: raw files are read-only; outputs are derived text + metadata."""
from pathlib import Path
import hashlib, json, csv, io, email
MAX_BYTES=5*1024*1024
SUPPORTED={'.txt','.csv','.json','.eml','.pdf','.xlsx'}

def ingest(path, seen_hashes=None):
    p=Path(path)
    if not p.exists() or not p.is_file(): raise ValueError('MISSING_INPUT')
    if p.suffix.lower() not in SUPPORTED: raise ValueError('UNSUPPORTED_TYPE')
    data=p.read_bytes()
    if not data: raise ValueError('EMPTY_INPUT')
    if len(data)>MAX_BYTES: raise ValueError('TOO_LARGE')
    sha=hashlib.sha256(data).hexdigest()
    if seen_hashes is not None and sha in seen_hashes: raise ValueError('DUPLICATE_INPUT')
    try:
        ext=p.suffix.lower()
        if ext=='.txt': text=data.decode('utf-8')
        elif ext=='.csv': text=data.decode('utf-8-sig')
        elif ext=='.json': text=json.dumps(json.loads(data.decode('utf-8')),ensure_ascii=False)
        elif ext=='.eml':
            msg=email.message_from_bytes(data)
            parts=[]
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type()=='text/plain': parts.append(part.get_payload(decode=True).decode(part.get_content_charset() or 'utf-8','replace'))
            else: parts.append(msg.get_payload(decode=True).decode(msg.get_content_charset() or 'utf-8','replace'))
            text='\n'.join(parts)
        elif ext=='.pdf':
            from pypdf import PdfReader
            r=PdfReader(io.BytesIO(data)); text='\n'.join((pg.extract_text() or '') for pg in r.pages)
        elif ext=='.xlsx':
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            lines=[]
            for ws in wb.worksheets:
                lines.append(f'[{ws.title}]')
                for row in ws.iter_rows(values_only=True): lines.append('\t'.join('' if v is None else str(v) for v in row))
            text='\n'.join(lines)
        else: raise ValueError('UNSUPPORTED_TYPE')
    except Exception as e: raise ValueError('PARSE_FAILED') from e
    if not text.strip(): raise ValueError('PARSE_FAILED')
    meta={'filename':p.name,'extension':p.suffix.lower(),'size_bytes':len(data),'sha256':sha,'derived_chars':len(text)}
    return {'text':text,'metadata':meta}
